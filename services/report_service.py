"""Monthly attendance reports and Excel export."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from database.db_manager import DatabaseManager
from services.staff_service import EmployeeRecord, StaffService
from services.work_time import (
    completed_day_minutes,
    day_presence_label,
    minutes_label,
    record_status_label,
    unique_latest_rows,
)
from utils.app_info import APP_NAME_FA, APP_VERSION
from utils.jalali_helper import JalaliHelper


@dataclass
class AttendanceDayRow:
    employee_id: int
    employee_name: str
    personnel_code: str
    jalali_date: str
    check_in: Optional[str]
    check_out: Optional[str]
    worked_minutes: Optional[int]
    delay_minutes: int
    early_leave_minutes: int
    overtime_minutes: int
    is_manual_entry: bool
    status: str
    status_label: str
    day_label: str
    is_incomplete: bool


@dataclass
class EmployeeMonthReport:
    employee: EmployeeRecord
    days: list[AttendanceDayRow]
    presence_days: int
    incomplete_days: int
    worked_minutes: int
    delay_minutes: int
    early_leave_minutes: int
    overtime_minutes: int


@dataclass
class MonthlyAttendanceReport:
    year: int
    month: int
    start_date: str
    end_date: str
    generated_at: str
    employees: list[EmployeeMonthReport] = field(default_factory=list)

    @property
    def month_label(self) -> str:
        return JalaliHelper.month_label(self.year, self.month)


class ReportService:
    def __init__(self, db: DatabaseManager) -> None:
        self.db = db
        self.staff = StaffService(db)

    def build_month_report(
        self,
        year: Optional[int] = None,
        month: Optional[int] = None,
        employee_id: Optional[int] = None,
        search: str = "",
    ) -> MonthlyAttendanceReport:
        today = JalaliHelper.today()
        year = year or today.year
        month = month or today.month
        start, end = JalaliHelper.month_bounds(year, month)
        generated_at = JalaliHelper.now_datetime_str()
        employees = self.staff.list_employees(include_inactive=True)
        if employee_id is not None:
            employees = [emp for emp in employees if emp.id == employee_id]
        query = (search or "").strip()
        if query:
            employees = [
                emp
                for emp in employees
                if query in emp.full_name or query in emp.personnel_code
                or query in emp.first_name
                or query in emp.last_name
            ]

        rows = self.db.list_attendances_in_range(start, end)
        by_employee: dict[int, list] = {}
        for row in rows:
            by_employee.setdefault(row["employee_id"], []).append(row)

        month_days = list(JalaliHelper.iter_dates(start, end))
        reports: list[EmployeeMonthReport] = []
        for employee in employees:
            unique_rows = unique_latest_rows(by_employee.get(employee.id, []))
            by_date = {row["jalali_date"]: row for row in unique_rows}
            day_rows: list[AttendanceDayRow] = []
            worked = delay = early = overtime = 0
            presence = incomplete = 0
            tolerance = employee.tolerance_minutes or 0
            for day in month_days:
                row = by_date.get(day)
                if row is None:
                    day_rows.append(
                        AttendanceDayRow(
                            employee_id=employee.id,
                            employee_name=employee.full_name,
                            personnel_code=employee.personnel_code,
                            jalali_date=day,
                            check_in=None,
                            check_out=None,
                            worked_minutes=None,
                            delay_minutes=0,
                            early_leave_minutes=0,
                            overtime_minutes=0,
                            is_manual_entry=False,
                            status="",
                            status_label="ثبت نشده",
                            day_label="عدم ثبت تردد",
                            is_incomplete=False,
                        )
                    )
                    continue
                minutes = completed_day_minutes(
                    row, employee.shift_start, employee.shift_end, tolerance
                )
                incomplete_flag = minutes is None
                if incomplete_flag:
                    incomplete += 1
                else:
                    presence += 1
                    worked += minutes or 0
                    delay += int(row["delay_minutes"] or 0)
                    early += int(row["early_leave_minutes"] or 0)
                    overtime += int(row["overtime_minutes"] or 0)
                day_rows.append(
                    AttendanceDayRow(
                        employee_id=employee.id,
                        employee_name=employee.full_name,
                        personnel_code=employee.personnel_code,
                        jalali_date=day,
                        check_in=row["check_in"],
                        check_out=row["check_out"],
                        worked_minutes=minutes,
                        delay_minutes=int(row["delay_minutes"] or 0),
                        early_leave_minutes=int(row["early_leave_minutes"] or 0),
                        overtime_minutes=int(row["overtime_minutes"] or 0),
                        is_manual_entry=bool(row["is_manual_entry"]),
                        status=row["status"] or "",
                        status_label=record_status_label(row),
                        day_label=day_presence_label(row),
                        is_incomplete=incomplete_flag,
                    )
                )
            reports.append(
                EmployeeMonthReport(
                    employee=employee,
                    days=day_rows,
                    presence_days=presence,
                    incomplete_days=incomplete,
                    worked_minutes=worked,
                    delay_minutes=delay,
                    early_leave_minutes=early,
                    overtime_minutes=overtime,
                )
            )
        reports.sort(key=lambda item: item.employee.full_name)
        return MonthlyAttendanceReport(
            year=year,
            month=month,
            start_date=start,
            end_date=end,
            generated_at=generated_at,
            employees=reports,
        )

    def default_excel_name(self, year: int, month: int) -> str:
        return f"attendance_report_{year}_{month:02d}.xlsx"

    def export_excel(self, report: MonthlyAttendanceReport, destination: str | Path) -> Path:
        path = Path(destination)
        try:
            path.parent.mkdir(parents=True, exist_ok=True)
            workbook = Workbook()
            self._write_summary_sheet(workbook.active, report)
            self._write_details_sheet(workbook.create_sheet("جزئیات تردد"), report)
            self._write_incomplete_sheet(workbook.create_sheet("ترددهای ناقص"), report)
            workbook.save(path)
        except PermissionError as exc:
            raise ValueError(
                "ذخیره فایل ممکن نشد. مسیر را بررسی کنید یا اگر فایل در برنامه دیگری باز است آن را ببندید."
            ) from exc
        except OSError as exc:
            raise ValueError(f"ذخیره فایل اکسل ممکن نشد: {exc}") from exc
        return path

    def _write_summary_sheet(self, sheet, report: MonthlyAttendanceReport) -> None:
        sheet.title = "خلاصه ماهانه"
        sheet.sheet_view.rightToLeft = True
        self._write_banner(sheet, report, "خلاصه ماهانه تردد پرسنل")
        headers = [
            "کد پرسنلی",
            "نام و نام خانوادگی",
            "وضعیت پرسنل",
            "تعداد روزهای حضور",
            "مجموع ساعات کارکرد",
            "مجموع تأخیر (دقیقه)",
            "مجموع تعجیل در خروج (دقیقه)",
            "مجموع اضافه‌کاری (دقیقه)",
            "تعداد ترددهای ناقص",
        ]
        start_row = 5
        self._write_header_row(sheet, start_row, headers)
        for index, item in enumerate(report.employees, start=start_row + 1):
            active = "فعال" if item.employee.is_active else "غیرفعال"
            sheet.append(
                [
                    item.employee.personnel_code,
                    item.employee.full_name,
                    active,
                    item.presence_days,
                    minutes_label(item.worked_minutes),
                    item.delay_minutes,
                    item.early_leave_minutes,
                    item.overtime_minutes,
                    item.incomplete_days,
                ]
            )
            self._style_data_row(sheet, index, len(headers))
        self._finalize_sheet(sheet, start_row, headers, max(len(report.employees), 1), "TblSummary")

    def _write_details_sheet(self, sheet, report: MonthlyAttendanceReport) -> None:
        sheet.sheet_view.rightToLeft = True
        self._write_banner(sheet, report, "جزئیات تردد روزانه")
        headers = [
            "کد پرسنلی",
            "نام پرسنل",
            "تاریخ شمسی",
            "ساعت ورود",
            "ساعت خروج",
            "ساعات کارکرد روزانه",
            "تأخیر (دقیقه)",
            "تعجیل در خروج (دقیقه)",
            "اضافه‌کاری (دقیقه)",
            "ثبت دستی",
            "وضعیت تردد",
            "وضعیت روز",
        ]
        start_row = 5
        self._write_header_row(sheet, start_row, headers)
        row_index = start_row + 1
        count = 0
        for item in report.employees:
            for day in item.days:
                if not day.check_in and not day.check_out:
                    continue
                sheet.append(
                    [
                        day.personnel_code,
                        day.employee_name,
                        day.jalali_date,
                        day.check_in or "",
                        day.check_out or "",
                        minutes_label(day.worked_minutes),
                        day.delay_minutes,
                        day.early_leave_minutes,
                        day.overtime_minutes,
                        "دستی" if day.is_manual_entry else "عادی",
                        day.status_label,
                        day.day_label,
                    ]
                )
                self._style_data_row(sheet, row_index, len(headers))
                row_index += 1
                count += 1
        self._finalize_sheet(sheet, start_row, headers, max(count, 1), "TblDetails")

    def _write_incomplete_sheet(self, sheet, report: MonthlyAttendanceReport) -> None:
        sheet.sheet_view.rightToLeft = True
        self._write_banner(sheet, report, "ترددهای ناقص نیازمند بررسی")
        headers = [
            "کد پرسنلی",
            "نام پرسنل",
            "تاریخ شمسی",
            "ساعت ورود",
            "ساعت خروج",
            "ثبت دستی",
            "وضعیت",
        ]
        start_row = 5
        self._write_header_row(sheet, start_row, headers)
        row_index = start_row + 1
        count = 0
        for item in report.employees:
            for day in item.days:
                if not day.is_incomplete:
                    continue
                sheet.append(
                    [
                        day.personnel_code,
                        day.employee_name,
                        day.jalali_date,
                        day.check_in or "",
                        day.check_out or "",
                        "دستی" if day.is_manual_entry else "عادی",
                        day.status_label,
                    ]
                )
                self._style_data_row(sheet, row_index, len(headers))
                row_index += 1
                count += 1
        self._finalize_sheet(sheet, start_row, headers, max(count, 1), "TblIncomplete")

    def _write_banner(self, sheet, report: MonthlyAttendanceReport, title: str) -> None:
        sheet.merge_cells("A1:I1")
        sheet.merge_cells("A2:I2")
        sheet.merge_cells("A3:I3")
        title_font = Font(name="Tahoma", size=16, bold=True, color="FFF8F1")
        meta_font = Font(name="Tahoma", size=11, color="3B2A1F")
        fill = PatternFill("solid", fgColor="5A3D2B")
        sheet["A1"] = f"{APP_NAME_FA} — {title}"
        sheet["A1"].font = title_font
        sheet["A1"].fill = fill
        sheet["A1"].alignment = Alignment(horizontal="right", vertical="center")
        for cell in sheet[1]:
            cell.fill = fill
        sheet["A2"] = f"ماه گزارش: {report.month_label}  |  بازه {report.start_date} تا {report.end_date}"
        sheet["A2"].font = meta_font
        sheet["A2"].alignment = Alignment(horizontal="right")
        sheet["A3"] = (
            f"تاریخ تولید: {report.generated_at}  |  نسخه برنامه {APP_VERSION}  |  "
            f"{datetime.now().strftime('%Y-%m-%d %H:%M')}"
        )
        sheet["A3"].font = meta_font
        sheet["A3"].alignment = Alignment(horizontal="right")
        sheet.row_dimensions[1].height = 24

    def _write_header_row(self, sheet, row: int, headers: list[str]) -> None:
        fill = PatternFill("solid", fgColor="6F4E37")
        font = Font(name="Tahoma", bold=True, color="FFF8F1")
        for col, header in enumerate(headers, start=1):
            cell = sheet.cell(row=row, column=col, value=header)
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def _style_data_row(self, sheet, row: int, columns: int) -> None:
        font = Font(name="Tahoma", size=11)
        align = Alignment(horizontal="center", vertical="center")
        thin = Border(
            left=Side(style="thin", color="E4D5C5"),
            right=Side(style="thin", color="E4D5C5"),
            top=Side(style="thin", color="E4D5C5"),
            bottom=Side(style="thin", color="E4D5C5"),
        )
        for col in range(1, columns + 1):
            cell = sheet.cell(row=row, column=col)
            cell.font = font
            cell.alignment = align
            cell.border = thin

    def _finalize_sheet(
        self, sheet, header_row: int, headers: list[str], data_rows: int, table_name: str
    ) -> None:
        last_row = header_row + data_rows
        last_col = get_column_letter(len(headers))
        ref = f"A{header_row}:{last_col}{last_row}"
        sheet.auto_filter.ref = ref
        sheet.freeze_panes = f"A{header_row + 1}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
        try:
            sheet.add_table(table)
        except ValueError:
            pass
        widths = [18, 28, 18, 16, 18, 20, 18, 22, 18, 14, 24, 18]
        for index, header in enumerate(headers, start=1):
            width = widths[index - 1] if index - 1 < len(widths) else max(14, len(header) + 4)
            sheet.column_dimensions[get_column_letter(index)].width = width
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
