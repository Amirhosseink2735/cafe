"""Phase 8 monthly attendance report and Excel export tests."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from database.db_manager import DatabaseManager
from services.backup_service import BackupService
from services.payroll_service import PayrollService
from services.report_service import ReportService
from services.staff_service import StaffService
from utils.jalali_helper import JalaliHelper


@pytest.fixture
def report_setup(tmp_path: Path):
    db = DatabaseManager(tmp_path / "cafe.db")
    db.initialize()
    staff = StaffService(db)
    employee_id = staff.save_employee("علی", "محمدی", "1001", "", "150000")
    staff.save_shift(employee_id, "08:00", "16:00", "5")
    inactive_id = staff.save_employee("سارا", "احمدی", "2002", "", "120000")
    staff.save_shift(inactive_id, "09:00", "17:00", "5")
    staff.deactivate_employee(inactive_id)
    reports = ReportService(db)
    payroll = PayrollService(db)
    return reports, payroll, staff, employee_id, inactive_id, db


def test_month_bounds_cover_full_jalali_month() -> None:
    start, end = JalaliHelper.month_bounds(1403, 12)
    days = list(JalaliHelper.iter_dates(start, end))
    assert start.endswith("/01")
    assert days[0] == start
    assert days[-1] == end
    assert len(days) in {29, 30}


def test_report_includes_inactive_and_incomplete(report_setup, tmp_path: Path) -> None:
    reports, payroll, _staff, employee_id, inactive_id, _db = report_setup
    today = JalaliHelper.today()
    start, end = JalaliHelper.month_bounds(today.year, today.month)
    days = list(JalaliHelper.iter_dates(start, min(JalaliHelper.today_str(), end)))
    work_day = days[-1]
    reports.db.insert_attendance(employee_id, work_day, "08:20:00", delay_minutes=20)
    reports.db.complete_attendance(
        reports.db.get_open_attendance(employee_id, work_day)["id"],
        "16:30:00",
        overtime_minutes=30,
    )
    if len(days) >= 2:
        missed = days[0] if days[0] != work_day else days[1]
        reports.db.insert_attendance(employee_id, missed, "08:00:00", status="missed_checkout")
    past = JalaliHelper.month_bounds(1403, 12)[1]
    reports.db.insert_manual_attendance(
        inactive_id, past, "09:00:00", "17:00:00", "completed", 0, 0, 0
    )

    current = reports.build_month_report(today.year, today.month)
    names = {item.employee.personnel_code: item for item in current.employees}
    assert "1001" in names
    assert "2002" in names
    ali = names["1001"]
    assert ali.presence_days == 1
    if len(days) >= 2:
        assert ali.incomplete_days == 1
        incomplete_rows = [day for day in ali.days if day.is_incomplete]
        assert incomplete_rows and incomplete_rows[0].worked_minutes is None
    pay = next(item for item in payroll.month_summaries() if item.employee_id == employee_id)
    assert ali.worked_minutes == pay.worked_minutes
    assert ali.delay_minutes == 20
    assert ali.overtime_minutes == 30

    esfand = reports.build_month_report(1403, 12)
    sara = next(item for item in esfand.employees if item.employee.id == inactive_id)
    assert sara.presence_days == 1
    assert sara.employee.is_active is False


def test_search_and_single_employee_filter(report_setup) -> None:
    reports, _payroll, _staff, employee_id, _inactive_id, _db = report_setup
    today = JalaliHelper.today()
    only_ali = reports.build_month_report(today.year, today.month, employee_id=employee_id)
    assert len(only_ali.employees) == 1
    assert only_ali.employees[0].employee.id == employee_id
    by_code = reports.build_month_report(today.year, today.month, search="1001")
    assert {item.employee.personnel_code for item in by_code.employees} == {"1001"}
    by_name = reports.build_month_report(today.year, today.month, search="سارا")
    assert {item.employee.personnel_code for item in by_name.employees} == {"2002"}


def test_excel_matches_on_screen_filters(report_setup, tmp_path: Path) -> None:
    reports, payroll, _staff, employee_id, _inactive_id, _db = report_setup
    today = JalaliHelper.today()
    start, end = JalaliHelper.month_bounds(today.year, today.month)
    work_day = min(JalaliHelper.today_str(), end)
    reports.db.insert_attendance(employee_id, work_day, "08:20:00", delay_minutes=20)
    reports.db.complete_attendance(
        reports.db.get_open_attendance(employee_id, work_day)["id"],
        "16:30:00",
        overtime_minutes=30,
    )
    report = reports.build_month_report(today.year, today.month, employee_id=employee_id)
    path = tmp_path / reports.default_excel_name(today.year, today.month)
    reports.export_excel(report, path)
    assert path.name == f"attendance_report_{today.year}_{today.month:02d}.xlsx"
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["خلاصه ماهانه", "جزئیات تردد", "ترددهای ناقص"]
    summary = workbook["خلاصه ماهانه"]
    assert summary.sheet_view.rightToLeft is True
    assert summary.freeze_panes == "A6"
    codes = [cell.value for cell in summary["A"] if cell.value not in (None, "کد پرسنلی")]
    assert "1001" in codes
    assert "2002" not in codes
    row = next(
        tuple(cell.value for cell in excel_row)
        for excel_row in summary.iter_rows(min_row=6, values_only=False)
        if excel_row[0].value == "1001"
    )
    payroll_row = next(item for item in payroll.month_summaries() if item.employee_id == employee_id)
    ali = report.employees[0]
    assert ali.worked_minutes == payroll_row.worked_minutes == 495
    assert any(cell == "8:15" for cell in row)


def test_excel_permission_error_message(report_setup, tmp_path: Path) -> None:
    reports, _payroll, _staff, employee_id, _inactive_id, _db = report_setup
    today = JalaliHelper.today()
    report = reports.build_month_report(today.year, today.month, employee_id=employee_id)
    blocked = tmp_path / "already_a_folder"
    blocked.mkdir()
    with pytest.raises(ValueError, match="ذخیره"):
        reports.export_excel(report, blocked)


def test_backup_and_restore_roundtrip(report_setup, tmp_path: Path) -> None:
    _reports, _payroll, staff, employee_id, _inactive_id, db = report_setup
    backups = BackupService(db)
    dest = tmp_path / "copy.db"
    backups.backup_to(dest)
    staff.save_employee("نیما", "رضایی", "3003", "", "90000")
    backups.restore_from(dest)
    codes = {emp.personnel_code for emp in staff.list_employees()}
    assert "3003" not in codes
    assert "1001" in codes
    with pytest.raises(ValueError):
        backups.restore_from(tmp_path / "missing.db")
    junk = tmp_path / "not_a_db.db"
    junk.write_text("hello", encoding="utf-8")
    with pytest.raises(ValueError, match="معتبر"):
        backups.restore_from(junk)
    stub = tmp_path / "partial.db"
    import sqlite3

    with sqlite3.connect(stub) as connection:
        connection.execute("CREATE TABLE employees (id INTEGER)")
        connection.execute("CREATE TABLE attendances (id INTEGER)")
        connection.execute("CREATE TABLE settings (key TEXT)")
        connection.commit()
    with pytest.raises(ValueError, match="برنامه"):
        backups.restore_from(stub)
