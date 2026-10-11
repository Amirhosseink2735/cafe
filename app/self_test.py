"""Headless checks for a frozen or installed CafeManagement build.

Does not open the interactive GUI loop. Honors CAFE_USER_DATA so a
packaged run can avoid the customer's live database.
"""

from __future__ import annotations

import traceback
from pathlib import Path

from database.db_manager import DatabaseManager
from gui.style import FONT_FAMILY, apply_theme, close_window, load_persian_fonts
from services.attendance_service import AttendanceService
from services.auth_service import AuthService
from services.report_service import ReportService
from services.staff_service import StaffService
from utils.jalali_helper import JalaliHelper
from utils.paths import icon_path, theme_path, user_data_dir


def run_self_test() -> int:
    data = user_data_dir()
    log_path = data / "self_test_result.txt"
    lines: list[str] = []
    try:
        theme = theme_path()
        icon = icon_path()
        if not theme.exists():
            raise FileNotFoundError(f"theme missing: {theme}")
        if not icon.exists():
            raise FileNotFoundError(f"icon missing: {icon}")
        lines.append(f"theme={theme}")
        lines.append(f"icon={icon}")

        apply_theme("light")
        family = load_persian_fonts()
        lines.append(f"font={family}")
        if family != "Vazirmatn":
            raise RuntimeError(f"expected Vazirmatn, got {family!r} ({FONT_FAMILY})")

        db = DatabaseManager(data / "cafe.db")
        db.initialize()
        auth = AuthService(db)
        if not auth.verify_password("admin"):
            raise RuntimeError("default admin password rejected")
        if auth.verify_password("not-the-password"):
            raise RuntimeError("wrong password accepted")
        lines.append("auth=ok")

        staff = StaffService(db)
        employee_id = staff.save_employee("آزمایش", "نصب", "9001", "", "100000")
        staff.save_shift(employee_id, "08:00", "16:00", "5")
        today = JalaliHelper.today_str()
        attendance = AttendanceService(db)
        attendance.check_in(employee_id, jalali_date=today, check_in_time="08:10:00")
        attendance.check_out(employee_id, jalali_date=today, check_out_time="16:05:00")
        lines.append("attendance=ok")

        reports = ReportService(db)
        today_j = JalaliHelper.today()
        report = reports.build_month_report(today_j.year, today_j.month, employee_id=employee_id)
        excel_path = data / "self_test_report.xlsx"
        saved = reports.export_excel(report, excel_path)
        from openpyxl import load_workbook

        workbook = load_workbook(saved)
        if "خلاصه ماهانه" not in workbook.sheetnames:
            raise RuntimeError("excel sheets missing")
        lines.append(f"excel={saved}")

        from gui.admin_window import AdminWindow
        from gui.kiosk_window import KioskWindow
        from gui.login_window import LoginWindow
        from gui.role_window import RoleWindow

        home = RoleWindow(db)
        home.withdraw()
        home.update_idletasks()
        close_window(home)

        login = LoginWindow(db)
        login.withdraw()
        login.password_entry.insert(0, "admin")
        login._login()
        if not login.authenticated:
            raise RuntimeError("login window rejected admin")

        admin = AdminWindow(db)
        admin.withdraw()
        admin.update_idletasks()
        admin.report_frame.refresh()
        close_window(admin)

        kiosk = KioskWindow(db)
        kiosk.withdraw()
        kiosk.update_idletasks()
        close_window(kiosk)
        lines.append("gui=ok")

        log_path.write_text("OK\n" + "\n".join(lines) + "\n", encoding="utf-8")
        return 0
    except Exception:
        log_path.parent.mkdir(parents=True, exist_ok=True)
        log_path.write_text("FAIL\n" + traceback.format_exc(), encoding="utf-8")
        return 1
